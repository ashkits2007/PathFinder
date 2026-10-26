"""Combined PATHFINDER agent module.

Provides both ``agent_app`` for Gemini/TinyFish repository research and
``process_documents_and_run_chain`` for local PDF, TXT, and PPTX analysis.
This is a copy candidate; it does not modify either source module.
"""

import os
from pathlib import Path
from typing import Literal, Sequence, TypedDict

import requests
from dotenv import load_dotenv
from langchain_chroma import Chroma
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_core.messages import HumanMessage
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.runnables import RunnablePassthrough
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_ollama import ChatOllama, OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langgraph.graph import END, StateGraph
from pptx import Presentation

load_dotenv()

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".pptx"}


class AgentState(TypedDict):
    query: str
    repo_link: str
    response_language: str
    domain: str
    fetched_markdown: str
    plan: str


def _make_llm() -> ChatGoogleGenerativeAI:
    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
        raise RuntimeError("Set GOOGLE_API_KEY in the environment before using /api/research.")
    return ChatGoogleGenerativeAI(
        model=os.getenv("GOOGLE_MODEL", "gemini-2.5-flash"),
        temperature=0,
        google_api_key=api_key,
    )


def invoke_llm(prompt: str) -> str:
    """Call Gemini and consistently return plain text from either response shape."""
    content = _make_llm().invoke([HumanMessage(content=prompt)]).content
    if isinstance(content, list):
        return "\n".join(
            str(part.get("text", part)) if isinstance(part, dict) else str(part)
            for part in content
        )
    return str(content)


def analyze_intent(state: AgentState):
    prompt = (
        "Classify the user's request as 'software' (applied codebase / implementation) "
        "or 'theoretical' (research synthesis). Reply with exactly one word: software "
        "or theoretical. Keep classification in English.\n\n"
        f"User request: {state.get('query', '')}\n"
        f"Repository: {state.get('repo_link', '') or 'not provided'}"
    )
    domain = invoke_llm(prompt).strip().lower()
    return {"domain": domain if domain in {"software", "theoretical"} else "software"}


def fetch_with_tinyfish(state: AgentState):
    repo_link = (state.get("repo_link") or "").strip()
    if not repo_link:
        return {"fetched_markdown": "No repository URL provided."}
    api_key = os.getenv("TINYFISH_API_KEY")
    if not api_key:
        return {"fetched_markdown": "Repository context was not fetched: TINYFISH_API_KEY is not configured."}
    try:
        response = requests.post(
            os.getenv("TINYFISH_API_URL", "https://api.fetch.tinyfish.ai"),
            json={"urls": [repo_link], "format": "markdown"},
            headers={"X-API-Key": api_key, "Content-Type": "application/json"},
            timeout=20,
        )
        response.raise_for_status()
        results = response.json().get("results", [])
        return {"fetched_markdown": results[0].get("text", "No text returned.") if results else "No content returned."}
    except requests.RequestException as exc:
        return {"fetched_markdown": f"Repository context could not be fetched ({type(exc).__name__})."}
    except (ValueError, KeyError, TypeError):
        return {"fetched_markdown": "Repository context response could not be parsed."}


def route_domain(state: AgentState) -> Literal["codebase_action_node", "literature_synthesis_node"]:
    return "codebase_action_node" if state["domain"] == "software" else "literature_synthesis_node"


def codebase_action_node(state: AgentState):
    language = state.get("response_language", "English")
    return {"plan": invoke_llm(
        f"User's goal/request: {state['query']}\nTarget repository: {state['repo_link']}\n"
        f"Repository context: {state['fetched_markdown']}\n\n"
        "Provide concrete implementation guidance and exact code changes where possible. "
        f"Write the response in {language}."
    )}


def literature_synthesis_node(state: AgentState):
    language = state.get("response_language", "English")
    return {"plan": invoke_llm(
        f"User's research goal: {state['query']}\nRetrieved context: {state['fetched_markdown']}\n\n"
        f"Synthesize the relevant theoretical concepts in {language}."
    )}


workflow = StateGraph(AgentState)
workflow.add_node("analyze_intent", analyze_intent)
workflow.add_node("fetch_with_tinyfish", fetch_with_tinyfish)
workflow.add_node("codebase_action_node", codebase_action_node)
workflow.add_node("literature_synthesis_node", literature_synthesis_node)
workflow.set_entry_point("analyze_intent")
workflow.add_edge("analyze_intent", "fetch_with_tinyfish")
workflow.add_conditional_edges("fetch_with_tinyfish", route_domain)
workflow.add_edge("codebase_action_node", END)
workflow.add_edge("literature_synthesis_node", END)
agent_app = workflow.compile()


def load_documents(file_paths: Sequence[str]) -> list[Document]:
    """Extract supported PDF, text, and PowerPoint files into LangChain documents."""
    documents: list[Document] = []
    for file_path in file_paths:
        path = Path(file_path)
        suffix = path.suffix.lower()
        if not path.is_file() or suffix not in ALLOWED_EXTENSIONS:
            continue
        if suffix == ".pdf":
            documents.extend(PyPDFLoader(str(path)).load())
        elif suffix == ".txt":
            text = path.read_text(encoding="utf-8", errors="replace").strip()
            if text:
                documents.append(Document(page_content=text, metadata={"source": path.name}))
        else:
            presentation = Presentation(str(path))
            for slide_number, slide in enumerate(presentation.slides, start=1):
                parts = [shape.text.strip() for shape in slide.shapes
                         if getattr(shape, "has_text_frame", False) and shape.text.strip()]
                if parts:
                    documents.append(Document(page_content="\n".join(parts), metadata={"source": path.name, "slide": slide_number}))
    return documents


def format_docs(documents: Sequence[Document]) -> str:
    return "\n\n".join(document.page_content for document in documents)


def process_documents_and_run_chain(
    query: str,
    repo_url: str | None = None,
    pdf_paths: Sequence[str] | None = None,
    response_language: str = "English",
    documents: Sequence[Document] | None = None,
) -> str:
    """Answer a literature query against uploaded documents using local Ollama."""
    full_query = query.strip()
    if repo_url:
        full_query += f"\n\nReference document URL: {repo_url}"
    source_documents = list(documents) if documents is not None else load_documents(pdf_paths or [])
    llm = ChatOllama(model="llama3.1", temperature=0.2)
    if not source_documents:
        return str(llm.invoke(f"Answer in {response_language}.\n\n{full_query[:12000]}").content)
    chunks = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120).split_documents(source_documents)
    vectorstore = Chroma.from_documents(documents=chunks, embedding=OllamaEmbeddings(model="all-minilm"))
    retriever = vectorstore.as_retriever(search_kwargs={"k": 4})
    prompt = ChatPromptTemplate.from_template(
        "Answer using the supplied context. State when it is insufficient. Respond in {language}.\n\n"
        "Context:\n{context}\n\nQuestion:\n{question}"
    )
    chain = (
        {"context": retriever | format_docs, "question": RunnablePassthrough(), "language": lambda _: response_language}
        | prompt | llm | StrOutputParser()
    )
    return chain.invoke(full_query[:12000])
